#!/usr/bin/env python3
"""Fetch the ISO planning-document benchmark sources for ``benchmark-corridor``.

Companion to ``scripts/data/fetch_aeo_electricity.py`` (the AEO2025 outlook
source). This lands the two ISO planning documents of the FC-5 benchmark
inventory (``docs/forecast-determination-rubric.md`` §6 items 3-4) that publish
their projection tables as a **directly downloadable XLSX**:

  * ``ERCOT_CDR_2025`` — ERCOT Capacity, Demand and Reserves report, Dec 2025
    vintage (§6 item 3): planned capacity additions by tech, firm peak load and
    reserve margin. Horizon ends at Summer 2030, so it anchors 2030 ONLY.
  * ``PJM_LOAD_2026`` — PJM 2026 Load Forecast Report tables (§6 item 4):
    summer peak load (Table B-1) and annual net energy (Table E-1), 2026-2046,
    so it anchors all three corridor years.

**Every value is machine-extracted from the published workbook — never typed.**
That is the point of this script: the raw README's "transcribe the table into a
unified CSV" path is a rule-5 fat-finger hazard for hundreds of cells, exactly
the hazard that kept the AEO source on a deterministic fetcher. Here the fetcher
downloads the primary document, reads the named sheet/row/column, and writes the
canonical unified CSV that
:func:`scripts.lib.benchmark_corridor.parse_unified_csv` already reads with no
per-source parsing code. Each emitted row carries its exact locator (sheet +
row label + column header) in ``source_page``.

**Context, never a fit target (CLAUDE.md rule 13).** Nothing in the model is
tuned toward a value this script writes; FC-5 gates the *explanation* of a
model-vs-benchmark divergence, never its size.

The downloaded workbooks are large (the CDR is ~13 MB) and are **gitignored**
with their SHA256 recorded, following the corpus convention in CLAUDE.md
("Cloning & session data"); the small extracted CSVs are what gets committed, so
the datatype still curates offline from committed raw.

Usage:
    python scripts/data/fetch_iso_planning_benchmarks.py
    python scripts/data/fetch_iso_planning_benchmarks.py --sources PJM_LOAD_2026
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import sys
import urllib.request
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[2]))

from scripts.lib import benchmark_corridor as bc  # noqa: E402
from scripts.lib.clean_io import paths  # noqa: E402

# Corridor target years (shared with the AEO source).
CORRIDOR_YEARS: tuple[int, ...] = (2030, 2035, 2040)

# --- ERCOT CDR (Dec 2025) ---------------------------------------------------
ERCOT_CDR_URL = (
    "https://www.ercot.com/files/docs/2025/12/19/"
    "CapacityDemandandReservesReport_December2025.xlsx"
)
ERCOT_CDR_DOC = (
    "ERCOT Capacity, Demand and Reserves Report, December 2025 "
    "(CapacityDemandandReservesReport_December2025.xlsx)"
)
# The CDR "Seasonal Summary" sheet is a matrix: three columns per season-year
# ([1] Peak Load Hour, [2] Peak Net Load Hour, [3] the difference). The corridor
# takes the PEAK LOAD HOUR column, the CDR's headline basis.
ERCOT_SHEET = "Seasonal Summary"
ERCOT_SEASON = "Summer"

# Row label prefix -> (quantity, tech). Labels are matched by prefix because the
# workbook's row labels carry long parenthetical qualifiers. Only rows whose
# meaning maps cleanly onto the canonical vocabulary are taken; the CDR's many
# adjustment lines (load-resource deductions, switchable capacity, …) are
# deliberately NOT intaken — they are accounting steps, not corridor quantities.
ERCOT_ROWS: tuple[tuple[str, str, str, str], ...] = (
    # (row-label prefix, quantity, tech, unit)
    ("Firm Peak Load", "peak_demand", "total", "MW"),
    ("Reserve Margin", "reserve_margin", "total", "fraction"),
    ("Total Capacity", "capacity", "total", "MW"),
)
# Planned Resource Additions block: label prefix -> canonical tech.
ERCOT_PLANNED: dict[str, str] = {
    "Thermal Resources": "gas",
    "Coastal Wind": "wind",
    "Panhandle Wind": "wind",
    "Other Wind": "wind",
    "Far West Solar": "solar",
    "West Solar": "solar",
    "Other Solar": "solar",
    "Energy Storage": "storage",
}
ERCOT_NOTE = (
    "ERCOT CDR Summer peak-load-hour column. Wind/solar/storage capacities are "
    "the CDR's PEAK-HOUR CONTRIBUTION (ELCC/seasonal-derated), NOT nameplate — "
    "so they are not directly comparable to AEO's nameplate GW without "
    "reconciliation (rule 11: misalignment documented, never silently mixed). "
    "The Dec-2025 CDR horizon ends at Summer 2030, so this source anchors 2030 only."
)

# --- PJM 2026 Load Forecast -------------------------------------------------
PJM_LOAD_URL = (
    "https://www.pjm.com/-/media/DotCom/planning/res-adeq/load-forecast/"
    "2026-load-report-tables.xlsx"
)
PJM_LOAD_DOC = "PJM 2026 Load Forecast Report tables (2026-load-report-tables.xlsx)"
# (sheet, table label, row label, quantity, unit)
PJM_TABLES: tuple[tuple[str, str, str, str, str], ...] = (
    ("Table B1", "Table B-1 Summer Peak Load (MW)", "PJM RTO", "peak_demand", "MW"),
    (
        "Table E1",
        "Table E-1 Annual Net Energy (GWh)",
        "PJM RTO",
        "energy_demand",
        "GWh",
    ),
)
PJM_NOTE = (
    "PJM RTO-total row of the 2026 Load Forecast Report. Summer peak is the "
    "coincident RTO summer peak (MW); energy is annual net energy (GWh). The "
    "PJM RTO footprint is the model's PJM zone set aggregated, so this is an "
    "ISO-total anchor, not a per-zone one."
)


def _download(url: str, dest: Path) -> Path:
    """Download ``url`` to ``dest`` (immutable raw), returning the path."""
    dest.parent.mkdir(parents=True, exist_ok=True)
    print(f"  downloading {url}")
    req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"})
    with urllib.request.urlopen(req, timeout=180) as fh:
        dest.write_bytes(fh.read())
    print(f"  wrote {dest} ({dest.stat().st_size:,} B)")
    return dest


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _write_unified(rows: list[dict], out_path: Path) -> Path:
    """Write canonical unified-CSV rows (the format the generic reader parses)."""
    out_path.parent.mkdir(parents=True, exist_ok=True)
    with out_path.open("w", newline="") as fh:
        w = csv.DictWriter(fh, fieldnames=list(bc.CANONICAL_COLUMNS))
        w.writeheader()
        for r in sorted(
            rows, key=lambda d: (d["target_year"], d["quantity"], d["tech"])
        ):
            w.writerow(r)
    print(f"  wrote {len(rows)} unified rows -> {out_path}")
    return out_path


def _label(cell) -> str:
    return "" if cell is None else str(cell).strip()


def fetch_ercot_cdr(raw_root: Path | None = None) -> Path:
    """Extract the ERCOT CDR Dec-2025 corridor rows into the unified CSV."""
    import openpyxl

    root = Path(raw_root) if raw_root is not None else paths.RAW_DIR
    out_dir = root / bc.DATATYPE / "ercot-cdr-2025"
    xlsx = _download(
        ERCOT_CDR_URL, out_dir / "CapacityDemandandReservesReport_December2025.xlsx"
    )
    sha = _sha256(xlsx)
    wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
    ws = wb[ERCOT_SHEET]
    grid = list(ws.iter_rows(values_only=True))
    season_hdr, year_hdr, tag_hdr = grid[1], grid[2], grid[5]

    # Column index of the "Peak Load Hour" ([n], the first of each season-year
    # triple) for each Summer year within the corridor years.
    year_col: dict[int, int] = {}
    for c in range(len(year_hdr)):
        if _label(season_hdr[c]) != ERCOT_SEASON:
            continue
        try:
            yr = int(str(year_hdr[c]).strip())
        except (TypeError, ValueError):
            continue
        if yr not in CORRIDOR_YEARS or not _label(tag_hdr[c]):
            continue
        year_col.setdefault(yr, c)  # first tagged col = Peak Load Hour
    if not year_col:
        raise RuntimeError(
            f"no {ERCOT_SEASON} corridor year found in {ERCOT_SHEET}; "
            "the CDR layout changed — refusing to guess"
        )
    print(f"  CDR corridor years found: {sorted(year_col)}")

    rows: list[dict] = []
    in_planned = False
    for r in grid:
        labels = [_label(x) for x in r[:4]]
        label = next((x for x in labels if x), "")
        if label.startswith("Planned Resource Additions"):
            in_planned = True
        elif label.startswith("Total Capacity"):
            in_planned = False

        specs: list[tuple[str, str, str]] = []
        for prefix, quantity, tech, unit in ERCOT_ROWS:
            if label.startswith(prefix):
                specs.append((quantity, tech, unit))
        planned_tech = None
        if in_planned:
            for prefix, tech in ERCOT_PLANNED.items():
                if label.startswith(prefix):
                    planned_tech = tech
                    break
        for yr, c in sorted(year_col.items()):
            val = r[c] if c < len(r) else None
            if not isinstance(val, (int, float)):
                continue
            for quantity, tech, unit in specs:
                rows.append(
                    _row(
                        "ERCOT_CDR_2025",
                        "ERCOT",
                        "ERCOT",
                        "2025-12",
                        "protocol",
                        yr,
                        quantity,
                        tech,
                        float(val),
                        unit,
                        ERCOT_CDR_DOC,
                        f"sheet={ERCOT_SHEET!r} row={label!r} col={ERCOT_SEASON} {yr} (peak load hour)",
                        ERCOT_NOTE,
                    )
                )
            if planned_tech is not None:
                rows.append(
                    _row(
                        "ERCOT_CDR_2025",
                        "ERCOT",
                        f"ERCOT planned additions: {label}",
                        "2025-12",
                        "protocol",
                        yr,
                        "capacity",
                        planned_tech,
                        float(val),
                        "MW",
                        ERCOT_CDR_DOC,
                        f"sheet={ERCOT_SHEET!r} row={label!r} (Planned Resource Additions) "
                        f"col={ERCOT_SEASON} {yr} (peak load hour)",
                        ERCOT_NOTE,
                    )
                )
    if not rows:
        raise RuntimeError("ERCOT CDR extraction produced no rows — layout changed")
    _write_unified(rows, out_dir / "ercot-cdr-2025.csv")
    (out_dir / "SHA256SUMS.txt").write_text(f"{sha}  {xlsx.name}\n")
    return out_dir / "ercot-cdr-2025.csv"


def fetch_pjm_load(raw_root: Path | None = None) -> Path:
    """Extract the PJM 2026 Load Forecast corridor rows into the unified CSV."""
    import openpyxl

    root = Path(raw_root) if raw_root is not None else paths.RAW_DIR
    out_dir = root / bc.DATATYPE / "pjm-load-2026"
    xlsx = _download(PJM_LOAD_URL, out_dir / "2026-load-report-tables.xlsx")
    sha = _sha256(xlsx)
    wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)

    rows: list[dict] = []
    for sheet, table_label, row_label, quantity, unit in PJM_TABLES:
        grid = list(wb[sheet].iter_rows(values_only=True))
        hdr_idx, year_col = None, {}
        for i, r in enumerate(grid[:12]):
            cols = {
                c: int(v)
                for c, v in enumerate(r)
                if isinstance(v, (int, float)) and 2020 < float(v) < 2060
            }
            if len(cols) >= 5:
                hdr_idx, year_col = i, cols
                break
        if hdr_idx is None:
            raise RuntimeError(f"{sheet}: no year header row found — layout changed")
        target = {y: c for c, y in year_col.items() if y in CORRIDOR_YEARS}
        missing = sorted(set(CORRIDOR_YEARS) - set(target))
        if missing:
            print(f"  [note] {sheet}: corridor years absent from horizon: {missing}")
        data_row = next(
            (r for r in grid if _label(r[0]) == row_label),
            None,
        )
        if data_row is None:
            raise RuntimeError(f"{sheet}: row {row_label!r} not found — layout changed")
        for yr, c in sorted(target.items()):
            val = data_row[c] if c < len(data_row) else None
            if not isinstance(val, (int, float)):
                continue
            rows.append(
                _row(
                    "PJM_LOAD_2026",
                    "PJM",
                    "PJM RTO",
                    "2026-01",
                    "reference",
                    yr,
                    quantity,
                    "total",
                    float(val),
                    unit,
                    PJM_LOAD_DOC,
                    f"sheet={sheet!r} ({table_label}) row={row_label!r} col={yr}",
                    PJM_NOTE,
                )
            )
    if not rows:
        raise RuntimeError("PJM load extraction produced no rows — layout changed")
    _write_unified(rows, out_dir / "pjm-load-2026.csv")
    (out_dir / "SHA256SUMS.txt").write_text(f"{sha}  {xlsx.name}\n")
    return out_dir / "pjm-load-2026.csv"


def _row(
    source: str,
    iso: str,
    region: str,
    vintage: str,
    scenario: str,
    target_year: int,
    quantity: str,
    tech: str,
    value: float,
    unit: str,
    source_doc: str,
    source_page: str,
    note: str,
) -> dict:
    """Build one canonical unified-CSV row."""
    return {
        "source": source,
        "iso": iso,
        "region": region,
        "vintage": vintage,
        "scenario": scenario,
        "target_year": target_year,
        "quantity": quantity,
        "tech": tech,
        "value": round(float(value), 6),
        "unit": unit,
        "source_doc": source_doc,
        "source_page": source_page,
        "note": note,
    }


FETCHERS = {
    "ERCOT_CDR_2025": fetch_ercot_cdr,
    "PJM_LOAD_2026": fetch_pjm_load,
}


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument(
        "--sources",
        nargs="*",
        default=None,
        choices=sorted(FETCHERS),
        help="subset to fetch (default: all fetchable ISO planning sources)",
    )
    args = ap.parse_args(argv)
    for name in args.sources or sorted(FETCHERS):
        print(f"=== {name} ===")
        FETCHERS[name]()
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
