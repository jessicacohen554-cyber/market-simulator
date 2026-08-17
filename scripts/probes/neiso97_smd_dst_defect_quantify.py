"""Quantify the NEISO SMD 2018-2023 DST-naive workbook defect from committed bytes.

neiso-97, 2026-08-17 -- Phase 1 of the audit row O8 charter
(``docs/audit/third-party-audit-2026-08.md`` section 8). neiso-96 found
(``results/calibration/ASSESSMENT-neiso96-h12026-intake-2026-08-15.md``
section 1.4) that the 2018-2023 SMD workbook vintage publishes a flat 24 rows
on every calendar day, so ``derive_actual_lmp._neiso_sheet_series``'s
positional clock -- correct for the true-23/25-row 2024-2025 vintage --
displaces the DST-transition windows of those years in the committed
``data/raw/_validation-source/actual_lmp_hourly_NEISO.parquet``.

This probe MEASURES the defect for the owner card; it repairs nothing and
writes nothing outside ``results/calibration/``:

1. **Vintage census** -- per year, workbook hub-sheet row counts per day
   (flat-24 vs true 23/25), and each year's DST transition dates.
2. **Alignment against the market's own published day** (the ISO-NE daily
   historical-report route, ``fetch_neiso_smd_zonal_lmp.py``, fetched to a
   scratch dir -- NOT committed): which re-placement of the flat-24 rows
   reproduces the published hours, measured over all 9 SMD locations x
   both markets, spring and fall separately.
3. **Committed-parquet diff** -- the corrected hub construction (same
   workbook bytes re-placed; truth-day overlay where the daily report
   exists) against the committed parquet: affected slots, displacement
   windows, magnitudes, fabricated-cell inventory, NaN accounting, and the
   internal control that every slot OUTSIDE the affected windows is
   bit-identical at the parquet's own float32 precision.

NO YEAR IS SOLVED, SCORED OR REGISTERED. Committed inputs + a published
actuals source only; the committed parquet is read, never written.

Usage:
    uv run python scripts/probes/neiso97_smd_dst_defect_quantify.py \
        --truth-dir <dir with NEISO_smd_zonal_lmp_<year>.csv for the DST days>
"""

from __future__ import annotations

import argparse
import csv
import json
import sys
from collections import defaultdict
from pathlib import Path

import numpy as np
import pandas as pd

REPO = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(REPO / "src"))
sys.path.insert(0, str(REPO))

import openpyxl  # noqa: E402

from scripts.data import derive_actual_lmp as dal  # noqa: E402

OUT = REPO / "results" / "calibration" / "_neiso97_smd_dst_defect_quantify.json"

#: Old (flat-24) workbook vintage years, per the neiso-96 measurement.
OLD_VINTAGE = (2018, 2019, 2020, 2021, 2022, 2023)
ALL_YEARS = (2018, 2019, 2020, 2021, 2022, 2023, 2024, 2025)

SHEET_BY_LOC = dal.NEISO_REPORT_LOCATIONS  # location_id -> SMD sheet name


def day_lengths(year: int) -> dict[str, int]:
    """Real local-day length in hours per calendar date, from the tz database."""
    days = pd.date_range(f"{year}-01-01", f"{year+1}-01-01", freq="D")
    starts = days.tz_localize(dal._EASTERN_TZ)
    hours = ((starts[1:] - starts[:-1]) / pd.Timedelta(hours=1)).astype(int)
    return {f"{d:%Y-%m-%d}": int(h) for d, h in zip(days[:-1], hours)}


def dst_dates(year: int) -> tuple[str, str]:
    """(spring-forward date, fall-back date) for ``year``."""
    lens = day_lengths(year)
    spring = [d for d, h in lens.items() if h == 23]
    fall = [d for d, h in lens.items() if h == 25]
    assert len(spring) == 1 and len(fall) == 1, (year, spring, fall)
    return spring[0], fall[0]


def workbook_days(year: int) -> dict[str, dict[str, list]]:
    """Raw workbook rows per sheet per date, in published order.

    Returns ``{sheet: {date: [(he_label, da, rt), ...]}}`` for the hub sheet
    plus every zone sheet, applying exactly the row filter
    ``dal._neiso_sheet_series`` applies.
    """
    path = dal.LMP_DIR / "NEISO" / f"{year}_smd_hourly.xlsx"
    needed = {dal.NEISO_HUB_SHEET, *(s for ss in dal.NEISO_ZONE_MAP.values() for s in ss)}
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    out: dict[str, dict[str, list]] = {}
    for sheet in wb.sheetnames:
        if sheet not in needed:
            continue
        rows = wb[sheet].iter_rows(values_only=True)
        next(rows, None)  # header
        per_date: dict[str, list] = defaultdict(list)
        for r in rows:
            if r is None or r[0] is None or r[1] is None:
                continue
            he = str(r[1]).strip()
            if not he[:2].isdigit():
                continue
            date = f"{pd.Timestamp(r[0]).normalize():%Y-%m-%d}"
            per_date[date].append((he, r[dal.NEISO_DA_COL], r[dal.NEISO_RT_COL]))
        out[sheet] = dict(per_date)
    wb.close()
    return out


def truth_days(truth_dir: Path, year: int) -> dict[str, dict[str, list]]:
    """Daily-report truth rows per sheet per date: ``{sheet: {date: [(he, da, rt), ...]}}``.

    Rows are ordered by the published within-day position (``seq``), the same
    positional quantity the workbook clock uses.
    """
    path = truth_dir / f"NEISO_smd_zonal_lmp_{year}.csv"
    if not path.exists():
        return {}
    acc: dict[str, dict[str, dict[int, tuple]]] = defaultdict(lambda: defaultdict(dict))
    with path.open(newline="") as f:
        for r in csv.DictReader(f):
            sheet = SHEET_BY_LOC.get(r["location_id"])
            if sheet is None:
                continue
            acc[sheet][r["date"]][int(r["seq"])] = (
                r["hour_ending"],
                float(r["da_lmp"]),
                float(r["rt_lmp"]),
            )
    return {
        sheet: {
            date: [by_seq[k] for k in sorted(by_seq)] for date, by_seq in dates.items()
        }
        for sheet, dates in acc.items()
    }


def f32(x) -> np.ndarray:
    """Cast to float32 -- the precision the committed parquet actually stores."""
    return np.asarray(x, dtype=np.float32)


def f32_eq(a, b) -> bool:
    """Float32 equality with NaN == NaN."""
    a, b = f32([a])[0], f32([b])[0]
    return (np.isnan(a) and np.isnan(b)) or a == b


def spring_alignment(wb_day: list, truth_day: list, kind_ix: int) -> dict:
    """Which re-placement of a flat-24 spring day reproduces the published 23 hours.

    Mapping PHANTOM (measured on 2023-03-12, all 9 sheets): the workbook
    fabricates the nonexistent 02:00 row at position 1 -- labeled "02", its
    value the MEAN of its two neighbors -- so published position p takes
    workbook row p for p < 1 and workbook row p+1 for p >= 1; workbook row 1
    is the phantom. Mapping TAIL: the 24 rows are the 23 real hours in order
    plus one trailing extra -- published position p takes workbook row p.
    """
    n_truth = len(truth_day)
    mism = {"phantom": 0, "tail": 0}
    for p in range(n_truth):
        t = truth_day[p][1 + kind_ix]
        w_phantom = wb_day[p if p < 1 else p + 1][1 + kind_ix]
        w_tail = wb_day[p][1 + kind_ix]
        mism["phantom"] += 0 if f32_eq(w_phantom, t) else 1
        mism["tail"] += 0 if f32_eq(w_tail, t) else 1
    # The phantom's provenance, measured: neighbor mean, rounded to cents.
    ph = float(wb_day[1][1 + kind_ix])
    nb_mean = (float(wb_day[0][1 + kind_ix]) + float(wb_day[2][1 + kind_ix])) / 2.0
    mism["phantom_is_neighbor_mean"] = abs(ph - nb_mean) <= 0.005 + 1e-9
    return mism


def fall_alignment(wb_day: list, truth_day: list, kind_ix: int) -> dict:
    """Verify the flat-24 fall-back day against the published 25 hours.

    Expected (neiso-96): workbook row 0 == published position 0; workbook
    row 1 == the MEAN of published positions 1 and 2 (the repeated hour's two
    instances); workbook rows 2..23 == published positions 3..24.
    """
    out = {"row0": 0, "mean_pair": 0, "shifted": 0}
    out["row0"] += 0 if f32_eq(wb_day[0][1 + kind_ix], truth_day[0][1 + kind_ix]) else 1
    pair_mean = (truth_day[1][1 + kind_ix] + truth_day[2][1 + kind_ix]) / 2.0
    # The workbook's published mean is rounded to cents; compare at a cent.
    if abs(float(wb_day[1][1 + kind_ix]) - pair_mean) > 0.005 + 1e-9:
        out["mean_pair"] += 1
    for k in range(2, 24):
        if not f32_eq(wb_day[k][1 + kind_ix], truth_day[k + 1][1 + kind_ix]):
            out["shifted"] += 1
    out["pair"] = [truth_day[1][1 + kind_ix], truth_day[2][1 + kind_ix]]
    out["workbook_mean"] = float(wb_day[1][1 + kind_ix])
    return out


def corrected_hub_dense(
    year: int, wb: dict[str, dict[str, list]], truth: dict[str, dict[str, list]]
) -> dict[str, np.ndarray]:
    """The corrected dense 8760 hub series (da & rt) for ``year``.

    Per operating day: a day whose workbook row count equals its real local
    length keeps the committed positional placement byte-for-byte; a
    defective day is taken WHOLESALE from the daily-report truth when the
    report exists, else (2018 spring, RT report unpublished) repaired by
    re-placement of the workbook's own rows under the measured PHANTOM
    mapping, the phantom row dropped, unrecoverable slots left NaN.
    """
    lens = day_lengths(year)
    hub = wb[dal.NEISO_HUB_SHEET]
    hub_truth = truth.get(dal.NEISO_HUB_SHEET, {})
    out = {}
    for kind_ix, kind in ((0, "da"), (1, "rt")):
        idx, vals = [], []
        for date, rows in hub.items():
            day_start = pd.Timestamp(date).tz_localize(dal._EASTERN_TZ).tz_convert("UTC")
            L = lens[date]
            if len(rows) == L:
                placed = [(k, rows[k][1 + kind_ix]) for k in range(len(rows))]
            elif date in hub_truth and len(hub_truth[date]) == L:
                placed = [
                    (k, hub_truth[date][k][1 + kind_ix]) for k in range(L)
                ]
            elif len(rows) == 24 and L == 23:  # spring, no truth: phantom mapping
                placed = [(p, rows[p if p < 1 else p + 1][1 + kind_ix]) for p in range(23)]
            elif len(rows) == 24 and L == 25:  # fall, no truth: pair unrecoverable
                placed = [(0, rows[0][1 + kind_ix])] + [
                    (k + 1, rows[k][1 + kind_ix]) for k in range(2, 24)
                ]
            else:
                raise AssertionError((year, date, len(rows), L))
            for pos, v in placed:
                idx.append(day_start + pd.Timedelta(hours=pos))
                vals.append(v)
        ser = pd.Series(
            pd.to_numeric(pd.Series(vals), errors="coerce").to_numpy(),
            index=pd.DatetimeIndex(idx),
        )
        out[kind] = dal._densify_std(ser, year, dal._STD_TZ["NEISO"])
    return out


def committed_hub_dense(year: int) -> dict[str, np.ndarray]:
    """The committed parquet's hub da/rt arrays for ``year``."""
    df = pd.read_parquet(
        dal.HOURLY_OUT / "actual_lmp_hourly_NEISO.parquet"
    )
    df = df[df["year"] == year].sort_values("hour")
    assert len(df) == dal._HOURS_PER_YEAR, (year, len(df))
    return {"da": df["da"].to_numpy(), "rt": df["rt"].to_numpy()}


def std_slot(year: int, date: str, pos: int) -> int:
    """Std-clock 8760 slot of ``date``'s ``pos``-th real hour."""
    inst = pd.Timestamp(date).tz_localize(dal._EASTERN_TZ).tz_convert(
        "UTC"
    ) + pd.Timedelta(hours=pos)
    return int(
        dal._std_hour_index(pd.DatetimeIndex([inst]), year, dal._STD_TZ["NEISO"])[0]
    )


def windows(slots: list[int]) -> list[list[int]]:
    """Contiguous [start, end] slot windows from a sorted slot list."""
    if not slots:
        return []
    out, s, prev = [], slots[0], slots[0]
    for x in slots[1:]:
        if x == prev + 1:
            prev = x
            continue
        out.append([s, prev])
        s = prev = x
    out.append([s, prev])
    return out


def main() -> int:
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument("--truth-dir", type=Path, required=True)
    args = p.parse_args()

    res: dict = {"census": {}, "alignment": {}, "parquet_diff": {}, "headline": {}}

    # -- 1. vintage census ---------------------------------------------------
    wb_cache: dict[int, dict] = {}
    for year in ALL_YEARS:
        wb = workbook_days(year)
        wb_cache[year] = wb
        hub = wb[dal.NEISO_HUB_SHEET]
        counts = defaultdict(int)
        for date, rows in hub.items():
            counts[len(rows)] += 1
        spring, fall = dst_dates(year)
        res["census"][year] = {
            "days_by_rowcount": dict(sorted(counts.items())),
            "spring": {"date": spring, "workbook_rows": len(hub[spring]), "true_hours": 23},
            "fall": {"date": fall, "workbook_rows": len(hub[fall]), "true_hours": 25},
            "sheets_agree": all(
                len(wb[sh].get(spring, [])) == len(hub[spring])
                and len(wb[sh].get(fall, [])) == len(hub[fall])
                for sh in wb
            ),
        }

    # -- 2. alignment against the published day ------------------------------
    for year in OLD_VINTAGE:
        truth = truth_days(args.truth_dir, year)
        spring, fall = dst_dates(year)
        rec: dict = {}
        for label, date, fn in (("spring", spring, spring_alignment), ("fall", fall, fall_alignment)):
            day_rec: dict = {"date": date, "truth_available": False}
            hub_truth = truth.get(dal.NEISO_HUB_SHEET, {})
            if date in hub_truth:
                day_rec["truth_available"] = True
                agg: dict[str, dict] = {}
                for kind_ix, kind in ((0, "da"), (1, "rt")):
                    per_sheet = {}
                    for sheet, days in truth.items():
                        if date not in days or date not in wb_cache[year].get(sheet, {}):
                            continue
                        per_sheet[sheet] = fn(
                            wb_cache[year][sheet][date], days[date], kind_ix
                        )
                    # aggregate mismatch counts over sheets
                    keys = {k for v in per_sheet.values() for k in v if isinstance(v[k], int)}
                    agg[kind] = {
                        k: int(sum(v[k] for v in per_sheet.values())) for k in sorted(keys)
                    }
                    if label == "fall":
                        agg[kind]["hub_pair"] = per_sheet[dal.NEISO_HUB_SHEET]["pair"]
                        agg[kind]["hub_workbook_mean"] = per_sheet[dal.NEISO_HUB_SHEET][
                            "workbook_mean"
                        ]
                    agg[kind]["n_sheets"] = len(per_sheet)
                day_rec["by_market"] = agg
            rec[label] = day_rec
        res["alignment"][year] = rec

    # -- 3. committed-parquet diff -------------------------------------------
    grand = {"cells_changed": 0, "max_abs_delta": 0.0, "max_affected_value": 0.0}
    for year in OLD_VINTAGE:
        truth = truth_days(args.truth_dir, year)
        corrected = corrected_hub_dense(year, wb_cache[year], truth)
        committed = committed_hub_dense(year)
        spring, fall = dst_dates(year)
        # Expected affected slot set, derived from the measured mechanism:
        #   spring: real positions 1..22 of the spring day (the neighbor-mean
        #           phantom at position 1, then every true value placed an
        #           hour late) + the next day's first slot (the spill mean);
        #   fall:   real positions 1..24 of the fall day (pair mean, early
        #           placements, the lost last hour).
        expected = sorted(
            {std_slot(year, spring, p) for p in range(1, 23)}
            | {std_slot(year, spring, 22) + 1}
            | {std_slot(year, fall, p) for p in range(1, 25)}
        )
        yr: dict = {}
        for kind in ("da", "rt"):
            a, b = f32(committed[kind]), f32(corrected[kind])
            both = ~(np.isnan(a) & np.isnan(b))
            diff = np.zeros(len(a), dtype=bool)
            diff[both] = ~(a[both] == b[both])
            # NaN<->value transitions count as differences:
            diff |= np.isnan(a) != np.isnan(b)
            changed = np.flatnonzero(diff)
            outside = sorted(set(changed.tolist()) - set(expected))
            vv = np.flatnonzero(diff & ~np.isnan(a) & ~np.isnan(b))
            deltas = np.abs(a[vv].astype(float) - b[vv].astype(float))
            yr[kind] = {
                "cells_changed": int(len(changed)),
                "windows": windows(changed.tolist()),
                "changed_outside_expected_mechanism": outside,
                "value_to_value_cells": int(len(vv)),
                "max_abs_delta": round(float(deltas.max()), 4) if len(deltas) else 0.0,
                "mean_abs_delta": round(float(deltas.mean()), 4) if len(deltas) else 0.0,
                "committed_nan": int(np.isnan(a).sum()),
                "corrected_nan": int(np.isnan(b).sum()),
                "nan_filled_slots": [int(s) for s in np.flatnonzero(np.isnan(a) & ~np.isnan(b))],
                "max_affected_committed": round(float(np.nanmax(a[changed])), 2)
                if len(changed)
                else None,
                "max_affected_corrected": round(float(np.nanmax(b[changed])), 2)
                if len(changed)
                else None,
            }
            grand["cells_changed"] += int(len(changed))
            if len(deltas):
                grand["max_abs_delta"] = max(grand["max_abs_delta"], float(deltas.max()))
            if len(changed):
                grand["max_affected_value"] = max(
                    grand["max_affected_value"],
                    float(np.nanmax(a[changed])),
                    float(np.nanmax(b[changed])),
                )
        # fabricated-cell inventory (hub, committed values)
        a_da, a_rt = f32(committed["da"]), f32(committed["rt"])
        s_ph = std_slot(year, spring, 1)
        s_sp = std_slot(year, spring, 22) + 1
        s_mn = std_slot(year, fall, 1)
        s_lost = std_slot(year, fall, 24)
        yr["fabricated_cells"] = {
            "spring_phantom_slot": s_ph,
            "spring_spill_mean_slot": s_sp,
            "fall_pair_mean_slot": s_mn,
            "fall_lost_hour_slot_nan": s_lost,
            "committed": {
                "da": [round(float(a_da[s]), 2) for s in (s_ph, s_sp, s_mn)],
                "rt": [round(float(a_rt[s]), 2) for s in (s_ph, s_sp, s_mn)],
            },
        }
        res["parquet_diff"][year] = yr

    res["headline"] = {
        "grand_total_hub_cells_changed": grand["cells_changed"],
        "grand_max_abs_delta_usd_mwh": round(grand["max_abs_delta"], 2),
        "grand_max_affected_value_usd_mwh": round(grand["max_affected_value"], 2),
        "tail_threshold_usd_mwh": 300.0,
        "any_affected_cell_at_or_above_tail": grand["max_affected_value"] >= 300.0,
        "in_sample_2023": {
            k: {
                "cells_changed": res["parquet_diff"][2023][k]["cells_changed"],
                "max_abs_delta": res["parquet_diff"][2023][k]["max_abs_delta"],
            }
            for k in ("da", "rt")
        },
    }

    OUT.write_text(json.dumps(res, indent=1, default=str) + "\n")
    print(json.dumps(res["headline"], indent=1))
    print(f"wrote {OUT}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
