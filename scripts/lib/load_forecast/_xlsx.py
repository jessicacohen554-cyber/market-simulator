"""Small shared helpers for the workbook-backed ``load-forecast`` ISO specs.

Every publisher that issues a machine-readable workbook is parsed through these
rather than through per-ISO copies, so the ISO modules carry only the layout
knowledge that is genuinely theirs.
"""

from __future__ import annotations

from pathlib import Path
from typing import Iterator

import openpyxl


def sheet_rows(path: Path, sheet: str) -> list[tuple]:
    """Return every row of ``sheet`` as a tuple of cell values.

    Read-only and values-only (formula results, not formulae), which is what a
    published forecast workbook needs and what keeps the read cheap.

    Args:
        path: workbook path.
        sheet: sheet name.

    Returns:
        List of row tuples, in sheet order.

    Raises:
        KeyError: if the workbook has no such sheet.
    """
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet not in wb.sheetnames:
            raise KeyError(f"{path.name}: no sheet {sheet!r} (has {wb.sheetnames})")
        return [tuple(r) for r in wb[sheet].iter_rows(values_only=True)]
    finally:
        wb.close()


def header_index(rows: list[tuple], *must_contain: str) -> int:
    """Index of the first row containing every one of ``must_contain``.

    Locating the header by content rather than by a hard-coded row number keeps
    the parser working when a publisher adds or removes a title line.

    Args:
        rows: rows as returned by :func:`sheet_rows`.
        must_contain: header labels that must all be present in the row.

    Returns:
        Row index of the header.

    Raises:
        ValueError: if no row contains all the labels.
    """
    wanted = {m.strip().lower() for m in must_contain}
    for i, row in enumerate(rows):
        cells = {str(c).strip().lower() for c in row if c is not None}
        if wanted <= cells:
            return i
    raise ValueError(f"no header row containing {sorted(wanted)}")


def numeric_records(
    rows: list[tuple], header_row: int, year_col: int = 0
) -> Iterator[tuple[int, dict[str, float]]]:
    """Yield ``(year, {column_label: value})`` for each numeric data row.

    Rows whose year cell is not an integer-valued number are skipped, which
    drops the title, note and blank rows publishers interleave with their data.

    Args:
        rows: rows as returned by :func:`sheet_rows`.
        header_row: index of the header row (see :func:`header_index`).
        year_col: column index holding the year.

    Yields:
        ``(year, values)`` per data row; ``values`` omits non-numeric cells.
    """
    labels = [
        (i, str(c).strip()) for i, c in enumerate(rows[header_row]) if c is not None
    ]
    for row in rows[header_row + 1 :]:
        if year_col >= len(row):
            continue
        cell = row[year_col]
        if not isinstance(cell, (int, float)) or isinstance(cell, bool):
            continue
        if float(cell) != int(cell):
            continue
        vals = {
            label: float(row[i])
            for i, label in labels
            if i != year_col
            and i < len(row)
            and isinstance(row[i], (int, float))
            and not isinstance(row[i], bool)
        }
        yield int(cell), vals
